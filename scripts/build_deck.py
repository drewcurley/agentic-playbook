#!/usr/bin/env python3
"""Generate the Agentic Starter overview deck (15-min talk) as a .pptx.

Design system: dark slate canvas, one teal accent + one amber highlight,
generous whitespace, shape-based diagrams instead of bullet walls. Run with the
project venv: .venv-deck/bin/python scripts/build_deck.py
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---- palette ---------------------------------------------------------------
INK      = RGBColor(0x0E, 0x14, 0x1B)   # near-black slate (background)
PANEL    = RGBColor(0x1B, 0x26, 0x32)   # raised panel
PANEL2   = RGBColor(0x24, 0x33, 0x42)   # lighter panel
TEAL     = RGBColor(0x2D, 0xD4, 0xBF)   # primary accent
TEALDK   = RGBColor(0x12, 0x6E, 0x67)
AMBER    = RGBColor(0xF5, 0xB7, 0x4D)   # highlight
WHITE    = RGBColor(0xF4, 0xF7, 0xFA)
MUTE     = RGBColor(0x9A, 0xAB, 0xB8)   # muted text
LINE     = RGBColor(0x33, 0x44, 0x52)
GREEN    = RGBColor(0x4A, 0xD2, 0x95)
RED      = RGBColor(0xE8, 0x6A, 0x6A)

FONT  = "Calibri"
FONTH = "Calibri"   # headings (Calibri ships everywhere; safe)

EMU_W, EMU_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width  = EMU_W
prs.slide_height = EMU_H
BLANK = prs.slide_layouts[6]

SLIDE_NO = 0   # incremented on every slide(); footer shows this

# ---- helpers ---------------------------------------------------------------
def slide():
    global SLIDE_NO
    SLIDE_NO += 1
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    bg.fill.solid(); bg.fill.fore_color.rgb = INK
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    sp = bg._element; sp.getparent().remove(sp); s.shapes._spTree.insert(2, sp)
    return s

def _set_font(run, size, color, bold=False, italic=False, name=FONT):
    run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
    run.font.name = name; run.font.color.rgb = color

def box(s, x, y, w, h, fill=None, line=None, line_w=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        radius=0.08):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    b.shadow.inherit = False
    if fill is None:
        b.fill.background()
    else:
        b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None:
        b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(line_w)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            b.adjustments[0] = radius
        except Exception:
            pass
    return b

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph is list of (txt,size,color,bold,italic,name)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(2); tf.margin_top = tf.margin_bottom = Pt(2)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space_after); p.line_spacing = line_spacing
        for seg in para:
            txt, size, color = seg[0], seg[1], seg[2]
            bold = seg[3] if len(seg) > 3 else False
            italic = seg[4] if len(seg) > 4 else False
            name = seg[5] if len(seg) > 5 else FONT
            r = p.add_run(); r.text = txt
            _set_font(r, size, color, bold, italic, name)
    return tb

def accent_bar(s, x, y, w=2.4, h=0.09):
    box(s, x, y, w, h, fill=TEAL, shape=MSO_SHAPE.RECTANGLE)

def kicker(s, txt, x=0.9, y=0.62):
    text(s, x, y, 9, 0.4, [[(txt.upper(), 13, TEAL, True)]])

def title(s, txt, x=0.9, y=0.95, size=34, w=11.5, color=WHITE):
    text(s, x, y, w, 1.2, [[(txt, size, color, True, False, FONTH)]])

def footer(s, label, accent=TEAL):
    text(s, 0.9, 7.02, 8, 0.35, [[(label, 10, MUTE)]])
    text(s, 11.6, 7.02, 1.3, 0.35, [[(f"{SLIDE_NO:02d} / {TOTAL:02d}", 10, MUTE)]], align=PP_ALIGN.RIGHT)

def chip(s, x, y, w, txt, fill=PANEL2, fg=WHITE, size=12, h=0.42, bold=True):
    box(s, x, y, w, h, fill=fill)
    text(s, x, y, w, h, [[(txt, size, fg, bold)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

TOTAL = 26   # asserted against SLIDE_NO at the end

def section(s, eyebrow, big, sub, accent=TEAL):
    """Full-bleed section divider."""
    box(s, 0, 0, 13.333, 7.5, fill=INK, shape=MSO_SHAPE.RECTANGLE)
    box(s, 10.2, -1.4, 4.6, 4.6, fill=PANEL, shape=MSO_SHAPE.OVAL)
    box(s, 11.6, 3.4, 3.4, 3.4, fill=TEALDK, shape=MSO_SHAPE.OVAL)
    accent_bar(s, 0.95, 2.7, w=2.0)
    text(s, 0.9, 2.9, 11.0, 0.5, [[(eyebrow.upper(), 15, accent, True)]])
    text(s, 0.9, 3.35, 11.6, 1.6, [[(big, 46, WHITE, True, False, FONTH)]])
    text(s, 0.9, 5.0, 10.8, 1.2, [[(sub, 17, MUTE)]], line_spacing=1.15)

# ===========================================================================
# 1 — TITLE
# ===========================================================================
s = slide()
box(s, 0, 0, 13.333, 7.5, fill=INK, shape=MSO_SHAPE.RECTANGLE)
# accent geometry
box(s, 9.8, -1.2, 5, 5, fill=PANEL, shape=MSO_SHAPE.OVAL)
box(s, 11.2, 3.6, 3.6, 3.6, fill=TEALDK, shape=MSO_SHAPE.OVAL)
accent_bar(s, 0.95, 2.35, w=2.0)
text(s, 0.9, 2.55, 11.5, 2.2, [
    [("Agentic Starter", 54, WHITE, True, False, FONTH)],
    [("A governed framework for AI-driven development", 22, TEAL, False)],
], space_after=10)
text(s, 0.9, 4.75, 11.0, 1.6, [
    [("Conventions, living memory, a 9-agent review team, and mechanical "
      "enforcement — so AI agents make changes you can trust.", 16, MUTE)],
], line_spacing=1.15)
text(s, 0.9, 6.5, 11.5, 0.5, [[("Clone → bootstrap → real work   •   15-minute overview", 13, AMBER, True)]])

# ===========================================================================
# 2 — THE PROBLEM
# ===========================================================================
s = slide()
kicker(s, "Why this exists"); title(s, "AI agents are powerful — and inconsistent")
prob = [
    ("Cold starts", "Every session begins blind. The agent re-derives architecture from grep, every time."),
    ("Invented facts", "Plausible-but-wrong assumptions that look authoritative — and ship."),
    ("No memory", "The same task two weeks apart yields two different designs."),
    ("Silent breakage", "A change to one route quietly breaks a caller nobody mapped."),
]
x0 = 0.9; cw = 5.7; ch = 1.75; gap = 0.35
for k,(h,d) in enumerate(prob):
    cx = x0 + (k % 2)*(cw+gap); cy = 2.05 + (k//2)*(ch+0.3)
    box(s, cx, cy, cw, ch, fill=PANEL)
    box(s, cx, cy, 0.10, ch, fill=RED, shape=MSO_SHAPE.RECTANGLE)
    text(s, cx+0.35, cy+0.22, cw-0.6, ch-0.4, [
        [(h, 18, WHITE, True)],
        [(d, 13.5, MUTE)],
    ], space_after=7, line_spacing=1.08)
text(s, 0.9, 6.35, 11.5, 0.5, [[("The antidote: a stable set of conventions every agent reads first, "
    "plus enforcement that holds when nobody is watching.", 14, TEAL, True, True)]])
footer(s, "The problem")

# ===========================================================================
# 3 — WHAT IT IS (3 pillars)
# ===========================================================================
s = slide()
kicker(s, "The solution in one picture"); title(s, "Three pillars, working together")
pillars = [
    ("LIVING MEMORY", "/docs as source of truth", "Architecture, contracts, decisions — complete enough that an agent reading only /docs can commit safely.", TEAL),
    ("THE AGENT TEAM", "9 specialists, 3 rounds", "Every non-trivial change is reviewed by domain experts with non-negotiable redlines + 7 strategic lenses.", AMBER),
    ("ENFORCEMENT", "hooks · CI · branch rules", "Mechanical gates so the conventions are not optional — even under deadline pressure.", GREEN),
]
cw = 3.75; gap = 0.45; x0 = 0.9; cy = 2.15; ch = 3.7
for k,(tag,h,d,col) in enumerate(pillars):
    cx = x0 + k*(cw+gap)
    box(s, cx, cy, cw, ch, fill=PANEL)
    box(s, cx, cy, cw, 0.12, fill=col, shape=MSO_SHAPE.RECTANGLE)
    text(s, cx+0.32, cy+0.42, cw-0.64, ch-0.7, [
        [(tag, 13, col, True)],
        [(h, 21, WHITE, True)],
        [("", 6, WHITE)],
        [(d, 14, MUTE)],
    ], space_after=8, line_spacing=1.12)
text(s, 0.9, 6.25, 11.5, 0.6, [[("Pull out any one pillar and the other two lose their teeth. "
    "The friction is the feature.", 14, WHITE, False, True)]])
footer(s, "What it is")

# ===========================================================================
# 4 — LIVING MEMORY
# ===========================================================================
s = slide()
kicker(s, "Pillar 1 — Living memory"); title(s, "/docs is the project's source of truth")
box(s, 0.9, 2.05, 11.55, 1.15, fill=PANEL2)
text(s, 1.2, 2.18, 11.0, 0.95, [
    [("The bar every doc must clear:", 13, AMBER, True)],
    [("“An agent who has read only /docs can make a commit confident it won’t "
      "break an existing route, data contract, integration, or behavior.”", 15.5, WHITE, False, True)],
], space_after=5, line_spacing=1.1)
# two columns: tier1 / protocol
box(s, 0.9, 3.5, 5.7, 2.95, fill=PANEL)
text(s, 1.2, 3.7, 5.2, 2.7, [
    [("TIER 1  — before any feature work", 13, TEAL, True)],
    [("architecture · contracts ★ · api · data-model", 13.5, WHITE, True)],
    [("stack · testing · deployment · ownership · onboarding", 13.5, WHITE)],
    [("", 5, WHITE)],
    [("contracts.md = the exhaustive, code-grounded inventory of "
      "everything an autonomous change could break.", 13, MUTE)],
], space_after=7, line_spacing=1.1)
box(s, 6.75, 3.5, 5.7, 2.95, fill=PANEL)
text(s, 7.05, 3.7, 5.2, 2.7, [
    [("GROUNDING & COMPLETENESS PROTOCOL", 13, AMBER, True)],
    [("✓ Grounded in real code (cite paths)", 13.5, WHITE)],
    [("✓ No invented facts — UNKNOWN, not a guess", 13.5, WHITE)],
    [("✓ Exhaustive where it's a contract (no sampling)", 13.5, WHITE)],
    [("✓ States invariants + blast radius", 13.5, WHITE)],
    [("✓ Verified against source by reviewers", 13.5, WHITE)],
], space_after=7, line_spacing=1.12)
footer(s, "Living memory")

# ===========================================================================
# 5 — THE 9-AGENT TEAM
# ===========================================================================
s = slide()
kicker(s, "Pillar 2 — The team"); title(s, "Nine specialists, each owning a domain")
agents = [
    ("Analyst","scope"),("Architect","security"),("Data Eng","schema"),
    ("Backend","app logic"),("Frontend","UI feasibility"),("UX","user-facing"),
    ("Test Eng","coverage"),("DevOps","deploy safety"),("Coordinator","process"),
]
cw=3.75; ch=1.18; gx=0.45; gy=0.3; x0=0.9; y0=2.1
for k,(name,win) in enumerate(agents):
    cx = x0 + (k%3)*(cw+gx); cy = y0 + (k//3)*(ch+gy)
    box(s, cx, cy, cw, ch, fill=PANEL)
    box(s, cx, cy, 0.10, ch, fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
    text(s, cx+0.3, cy+0.16, cw-0.5, ch-0.3, [
        [(name, 16, WHITE, True)],
        [("wins on  ", 11, MUTE),(win, 11, AMBER, True)],
    ], space_after=3, line_spacing=1.0)
text(s, 0.9, 6.5, 11.5, 0.4, [[("Each carries non-negotiable redlines. Conflicts resolve by domain: "
    "Architect wins security, Test Eng wins coverage, Analyst wins scope.", 13, MUTE, False, True)]])
footer(s, "The agent team")

# ===========================================================================
# 6 — REVIEW CYCLE + LENSES
# ===========================================================================
s = slide()
kicker(s, "How review works"); title(s, "Three rounds, then seven lenses for big calls")
rounds = [
    ("ROUND 1 · Analysis","Analyst · Architect · Data Eng","scope, security, schema"),
    ("ROUND 2 · Implementation","Backend · Frontend · UX","feasibility, user impact"),
    ("ROUND 3 · Verification","Test Eng · DevOps","coverage, deploy safety"),
]
y=2.1; rh=1.15
for k,(h,who,what) in enumerate(rounds):
    cy=y+k*(rh+0.18)
    box(s, 0.9, cy, 7.6, rh, fill=PANEL)
    box(s, 0.9, cy, 1.5, rh, fill=TEALDK)
    text(s, 0.9, cy, 1.5, rh, [[(f"{k+1}", 30, TEAL, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 2.6, cy+0.16, 5.7, rh-0.3, [
        [(h, 15.5, WHITE, True)],
        [(who+"   ", 12, MUTE),("→ "+what, 12, AMBER)],
    ], space_after=3, line_spacing=1.0)
    if k<2:
        text(s, 4.4, cy+rh-0.04, 1, 0.3, [[("▼", 11, LINE)]], align=PP_ALIGN.CENTER)
# lenses panel
box(s, 8.8, 2.1, 3.65, 3.97, fill=PANEL)
text(s, 9.1, 2.3, 3.1, 3.7, [
    [("7 STRATEGIC LENSES", 13, AMBER, True)],
    [("for every major decision", 11, MUTE, False, True)],
    [("", 5, WHITE)],
    [("CEO · Purchasing · PM", 13, WHITE)],
    [("Adopter · Builder", 13, WHITE)],
    [("Investor · Marketing", 13, WHITE)],
    [("", 5, WHITE)],
    [("Surface conflicts between lenses — don't paper over them.", 12, MUTE)],
], space_after=6, line_spacing=1.1)
text(s, 0.9, 6.5, 11.5, 0.4, [[("Each round's blockers are resolved before the next begins. "
    "Output: a documented review artifact with all sign-offs.", 13, MUTE, False, True)]])
footer(s, "Review cycle")

# ===========================================================================
# 7 — ENFORCEMENT
# ===========================================================================
s = slide()
kicker(s, "Pillar 3 — Enforcement"); title(s, "Conventions that aren't optional")
layers = [
    ("LOCAL  ·  git hooks","No commits to main · no placeholders · secret scan · Task-ID required · /docs consensus assertion","Runs on every commit & push"),
    ("SERVER  ·  CI mirror","Same checks run server-side (Azure DevOps now; GitHub Actions parked) — shared scripts/ci-checks.sh","--no-verify can't bypass this"),
    ("BRANCH  ·  protection","PR required · no force-push · admins included · CI check required · human approval (team)","The keystone — makes the rest real"),
]
y=2.1; rh=1.3
for k,(h,d,tag) in enumerate(layers):
    cy=y+k*(rh+0.18)
    box(s, 0.9, cy, 11.55, rh, fill=PANEL)
    box(s, 0.9, cy, 0.12, rh, fill=GREEN, shape=MSO_SHAPE.RECTANGLE)
    text(s, 1.25, cy+0.16, 8.7, rh-0.3, [
        [(h, 15.5, WHITE, True)],
        [(d, 12.5, MUTE)],
    ], space_after=4, line_spacing=1.05)
    text(s, 10.0, cy, 2.4, rh, [[(tag, 11.5, AMBER, True)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s, "Enforcement")

# ===========================================================================
# 8 — TWO SHAPES (single vs multi-repo)
# ===========================================================================
s = slide()
kicker(s, "Two ways to adopt"); title(s, "Single repo, or a multi-repo workspace")
# single
box(s, 0.9, 2.1, 5.7, 4.0, fill=PANEL)
text(s, 1.2, 2.3, 5.1, 0.6, [[("SINGLE REPO", 15, TEAL, True)],[("the fork is your project", 12, MUTE)]], space_after=2)
box(s, 1.3, 3.35, 4.9, 2.5, fill=PANEL2)
text(s, 1.55, 3.55, 4.5, 2.2, [
    [("repo/", 14, WHITE, True)],
    [("  AGENTS.md · .githooks/", 12.5, MUTE)],
    [("  /docs  (living memory)", 12.5, WHITE)],
    [("  /tasks  · src/ ...", 12.5, MUTE)],
], space_after=6, line_spacing=1.1)
# multi
box(s, 6.75, 2.1, 5.7, 4.0, fill=PANEL)
text(s, 7.05, 2.3, 5.1, 0.6, [[("MULTI-REPO WORKSPACE", 15, AMBER, True)],[("one product, several repos — no monorepo", 12, MUTE)]], space_after=2)
box(s, 7.15, 3.35, 4.9, 2.5, fill=PANEL2)
text(s, 7.4, 3.5, 4.5, 2.3, [
    [("playbook/   ← shared, governs all", 12.5, WHITE, True)],
    [("  integration-map.md  ★ cross-repo seams", 12, TEAL)],
    [("web/   api/   infra/   (siblings)", 12.5, WHITE)],
    [("  each: own /docs · hooks · CI", 12, MUTE)],
], space_after=6, line_spacing=1.12)
text(s, 0.9, 6.35, 11.5, 0.5, [[("Multi-repo superpower: change a contract in one repo and the agent sees "
    "which sibling repos consume it — before it breaks them.", 13.5, WHITE, False, True)]])
footer(s, "Single vs. multi-repo")

# ===========================================================================
# 9 — SAFE AUTONOMOUS COMMITS (contracts loop)
# ===========================================================================
s = slide()
kicker(s, "The payoff"); title(s, "Why a change can be trusted")
steps = [
    ("1","Read","Agent reads /docs + contracts.md at session start"),
    ("2","Locate","Finds every route/schema/event its change touches"),
    ("3","Assess","Checks blast radius — incl. cross-repo consumers"),
    ("4","Prove","Runs the regression-safety check + contract tests"),
    ("5","Record","Updates contracts.md in the same PR — atomic"),
]
cw=2.18; gap=0.18; x0=0.9; cy=2.4; ch=2.6
for k,(num,h,d) in enumerate(steps):
    cx=x0+k*(cw+gap)
    box(s, cx, cy, cw, ch, fill=PANEL)
    box(s, cx, cy, cw, 0.7, fill=TEALDK)
    text(s, cx, cy+0.06, cw, 0.6, [[(num, 26, TEAL, True)]], align=PP_ALIGN.CENTER)
    text(s, cx+0.18, cy+0.85, cw-0.36, ch-1.0, [
        [(h, 15, WHITE, True)],
        [(d, 11.5, MUTE)],
    ], space_after=5, line_spacing=1.08)
    if k<4:
        text(s, cx+cw-0.06, cy, 0.3, ch, [[("›", 22, LINE)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 5.5, 11.5, 0.8, [[("Hard gate: ", 14, AMBER, True),
    ("no commit touching a documented contract merges without updating contracts.md "
     "in the same PR. Living memory can't lag the code.", 14, WHITE)]], line_spacing=1.1)
footer(s, "Safe commits")

# ===========================================================================
# 10 — THE JOURNEY (clone -> ready) overview
# ===========================================================================
s = slide()
kicker(s, "Clone → ready"); title(s, "Six steps from fork to enforced")
journey = [
    ("Fork / copy","Use as template; keep the layout"),
    ("Replace placeholders","grep -rl '[insert' — stack, handles, names"),
    ("Install hooks","setup-hooks.sh  (or setup-workspace.sh)"),
    ("Wire CI","Azure DevOps pipeline + PR build validation"),
    ("Protect main","PR required · approvals to match headcount"),
    ("First agent session","Agent bootstraps /docs before any feature work"),
]
y0=2.1; rh=0.66;
for k,(h,d) in enumerate(journey):
    cy=y0+k*(rh+0.12)
    box(s, 0.9, cy, 0.66, rh, fill=TEAL, shape=MSO_SHAPE.OVAL)
    text(s, 0.9, cy, 0.66, rh, [[(str(k+1), 19, INK, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    box(s, 1.8, cy, 10.65, rh, fill=PANEL)
    text(s, 2.1, cy, 4.6, rh, [[(h, 15, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 6.6, cy, 5.7, rh, [[(d, 12.5, MUTE)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 6.95, 11.5, 0.4, [[("One runbook: ", 12, MUTE, False, True),("ADOPTING.md", 12, AMBER, True, True),
    ("  — ordered, with don't-brick-the-repo warnings.", 12, MUTE, False, True)]])
footer(s, "The journey")

# ===========================================================================
# 11 — EXISTING PROJECT
# ===========================================================================
s = slide()
kicker(s, "Path A"); title(s, "Bootstrapping an existing codebase")
box(s, 0.9, 2.1, 11.55, 1.0, fill=PANEL2)
text(s, 1.2, 2.22, 11.0, 0.8, [[("The agent’s first job is to ", 14, WHITE),
    ("build /docs from the real code", 14, TEAL, True),
    (" — not feature work. It reads the repo and drafts living memory grounded in what exists.", 14, WHITE)]], line_spacing=1.1)
pts = [
    ("Map what's real","Routes, schema, integrations, env vars enumerated from source — not guessed. contracts.md is exhaustive."),
    ("Mark the unknowns","Anything unconfirmed becomes UNKNOWN — flagged for you, never invented."),
    ("Review to land","All 9 agents sign off before /docs lands; reviewers re-check claims against code."),
    ("Then feature work","Once Tier 1 is in, the agent can change code knowing what it might break."),
]
for k,(h,d) in enumerate(pts):
    cx=0.9+(k%2)*5.9; cy=3.35+(k//2)*1.55
    box(s, cx, cy, 5.65, 1.4, fill=PANEL)
    box(s, cx, cy, 0.1, 1.4, fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
    text(s, cx+0.32, cy+0.18, 5.2, 1.1, [[(h,15,WHITE,True)],[(d,12,MUTE)]], space_after=5, line_spacing=1.05)
footer(s, "Existing project")

# ===========================================================================
# 12 — NEW PROJECT
# ===========================================================================
s = slide()
kicker(s, "Path B"); title(s, "Starting a greenfield project")
box(s, 0.9, 2.1, 11.55, 1.0, fill=PANEL2)
text(s, 1.2, 2.22, 11.0, 0.8, [[("With little code yet, /docs is ", 14, WHITE),
    ("forward-looking", 14, AMBER, True),
    (" — it describes the intended state, clearly labeled as such (never blurred with “what exists”).", 14, WHITE)]], line_spacing=1.1)
pts = [
    ("Decide the stack","ADR 0001 captures the foundational choice + why — the question a future engineer asks."),
    ("Design contracts first","api.md / data-model.md describe planned routes & schema as the contract to build to."),
    ("Lenses up front","Major direction calls run the 7 lenses before code, while changing course is cheap."),
    ("Grow memory as you build","Each feature fills contracts.md with real, grounded entries — forward-looking → descriptive."),
]
for k,(h,d) in enumerate(pts):
    cx=0.9+(k%2)*5.9; cy=3.35+(k//2)*1.55
    box(s, cx, cy, 5.65, 1.4, fill=PANEL)
    box(s, cx, cy, 0.1, 1.4, fill=AMBER, shape=MSO_SHAPE.RECTANGLE)
    text(s, cx+0.32, cy+0.18, 5.2, 1.1, [[(h,15,WHITE,True)],[(d,12,MUTE)]], space_after=5, line_spacing=1.05)
footer(s, "New project")

# ===========================================================================
# 13 — A DAY IN THE LIFE
# ===========================================================================
s = slide()
kicker(s, "In practice"); title(s, "A typical session, end to end")
flow = [
    ("Session start","Agent reads /docs + tasks; one-line ack of what it loaded"),
    ("Open a task","tasks/T-NNN.md — acceptance criteria, branch off main"),
    ("Plan → review","3-round review on the plan before a line of code"),
    ("Build → review","Implement, then review the build; tests for every path"),
    ("Regression-safety","Diff vs contracts.md; update docs in the same PR"),
    ("Commit → PR","Task-ID enforced; state → completed in the work PR"),
]
for k,(h,d) in enumerate(flow):
    cx=0.9+(k%2)*5.9; cy=2.15+(k//2)*1.35
    box(s, cx, cy, 5.65, 1.2, fill=PANEL)
    box(s, cx, cy, 0.55, 1.2, fill=TEALDK)
    text(s, cx, cy, 0.55, 1.2, [[(str(k+1),18,TEAL,True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, cx+0.75, cy+0.14, 4.7, 0.95, [[(h,14.5,WHITE,True)],[(d,11.5,MUTE)]], space_after=3, line_spacing=1.05)
text(s, 0.9, 6.45, 11.5, 0.5, [[("Heavy on the first PR, natural by the third. ", 13, WHITE, False, True),
    ("The ledger update rides in the work PR — no follow-up “mark it done” noise.", 13, MUTE, False, True)]])
footer(s, "A day in the life")

# ===========================================================================
# 14 — SCALING DOWN / SOLO
# ===========================================================================
s = slide()
kicker(s, "Right-sizing"); title(s, "Scales from solo to enterprise")
cols = [
    ("ONE HUMAN, MANY HATS","One person signs off as each specialist in turn — the role-switch forces the perspective a missing teammate would.", TEAL),
    ("SOLO-SAFE DEFAULTS","GitHub blocks self-approval, so a 1-person repo runs 0 required approvals; raise to ≥1 + code-owner review when a teammate joins.", AMBER),
    ("NEVER TURN OFF","Redlines, the hooks + CI mirror, and /docs consensus — calibrated to worst-case stakes for a reason.", GREEN),
]
cw=3.75; gap=0.45
for k,(h,d,col) in enumerate(cols):
    cx=0.9+k*(cw+gap); cy=2.15; ch=3.7
    box(s, cx, cy, cw, ch, fill=PANEL)
    box(s, cx, cy, cw, 0.12, fill=col, shape=MSO_SHAPE.RECTANGLE)
    text(s, cx+0.3, cy+0.4, cw-0.6, ch-0.7, [[(h,14.5,col,True)],[("",6,WHITE)],[(d,13.5,MUTE)]],
         space_after=8, line_spacing=1.15)
text(s, 0.9, 6.25, 11.5, 0.6, [[("Lower stakes? Scale the procedure down and ", 13.5, WHITE, False, True),
    ("write down what you removed and why", 13.5, AMBER, True, True),
    (" — don't pretend the full process is running when it isn't.", 13.5, WHITE, False, True)]])
footer(s, "Scaling down")

# ===========================================================================
# ACT II — CASE STUDY (divider)
# ===========================================================================
s = slide()
section(s, "Case study — a real production repo",
        "From “describes it” to “safe to act on”",
        "REFOPS · Living-Memory Hardening (T-002). Reviewed by 9 agents + 7 lenses. "
        "The codebase was never touched — only the living memory changed.",
        accent=AMBER)

# ---- CS1: the problem on a real repo --------------------------------------
s = slide()
kicker(s, "Case study · the gap"); title(s, "Docs that only describe aren't safe to act on")
box(s, 0.9, 2.05, 11.55, 1.05, fill=PANEL2)
text(s, 1.2, 2.18, 11.0, 0.9, [[("~1,100 lines of prose describing ~290 files of code = a lossy summary. "
    "A summary is wrong in exactly the places that cause outages.", 15, WHITE, False, True)]], line_spacing=1.12)
text(s, 0.9, 3.45, 11.5, 0.5, [[("Before hardening, /docs answered “what is this system?” — but not the three "
    "questions an autonomous agent must answer before it changes anything:", 14, MUTE)]], line_spacing=1.1)
q = [
    ("1","What contract must I NOT break?"),
    ("2","What else is coupled to this change?"),
    ("3","How do I PROVE I didn't break it?"),
]
for k,(num,t) in enumerate(q):
    cy=4.25+k*0.78
    box(s, 0.9, cy, 11.55, 0.66, fill=PANEL)
    box(s, 0.9, cy, 0.66, 0.66, fill=RED, shape=MSO_SHAPE.OVAL)
    text(s, 0.9, cy, 0.66, 0.66, [[(num, 18, WHITE, True)]], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.8, cy, 10.4, 0.66, [[(t, 15.5, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s, "Case study · the gap")

# ---- CS2: the bar — supervised autonomy -----------------------------------
s = slide()
kicker(s, "Case study · the bar"); title(s, "Supervised autonomy")
box(s, 0.9, 2.1, 11.55, 1.15, fill=PANEL2)
text(s, 1.2, 2.26, 11.0, 0.95, [[("Agent proposes + self-reviews against the docs.  ", 17, WHITE, True),
    ("A human gates every commit.", 17, AMBER, True)]], line_spacing=1.1)
pts=[
    ("Agent does the heavy lifting","…and shows its work — a complete, reviewable proposal."),
    ("Human approves mechanically","With a precise checklist in hand, not a vibe check."),
    ("No self-approved risk","No agent merges a risky change on its own authority."),
]
for k,(h,d) in enumerate(pts):
    cy=3.55+k*1.05
    box(s, 0.9, cy, 11.55, 0.92, fill=PANEL)
    box(s, 0.9, cy, 0.1, 0.92, fill=TEAL, shape=MSO_SHAPE.RECTANGLE)
    text(s, 1.25, cy+0.13, 11.0, 0.7, [[(h+"   ", 15, WHITE, True),(d, 13, MUTE)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 6.55, 11.5, 0.4, [[("The docs' job: let the agent hand the human a trustworthy "
    "“here's what I changed, here's what could break, here's the proof.”", 13, MUTE, False, True)]])
footer(s, "Case study · the bar")

# ---- CS3: the 4-step self-review gate -------------------------------------
s = slide()
kicker(s, "Case study · the core mechanism"); title(s, "A 4-step self-review gate")
steps=[
    ("1","CONTRACT","contracts.md","what must NOT break?", TEAL),
    ("2","BLAST RADIUS","invariants.md · features-map.md","what else must change?", AMBER),
    ("3","PROOF","verification.md","what check proves it's safe?", GREEN),
    ("4","HUMAN GATE","diff + proof + checklist","hand it to a person", RED),
]
cw=2.78; gap=0.2; x0=0.9; cy=2.3; ch=2.7
for k,(num,h,doc,q,col) in enumerate(steps):
    cx=x0+k*(cw+gap)
    box(s, cx, cy, cw, ch, fill=PANEL)
    box(s, cx, cy, cw, 0.66, fill=col if col!=AMBER else TEALDK)
    text(s, cx, cy+0.05, cw, 0.6, [[(num, 24, WHITE, True)]], align=PP_ALIGN.CENTER)
    text(s, cx+0.2, cy+0.82, cw-0.4, ch-1.0, [
        [(h, 14.5, col, True)],
        [(doc, 12, WHITE, True)],
        [("", 4, WHITE)],
        [(q, 12.5, MUTE)],
    ], space_after=5, line_spacing=1.08)
    if k<3: text(s, cx+cw-0.04, cy, 0.28, ch, [[("›", 20, LINE)]], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.9, 5.35, 11.5, 0.8, [[("A change is not “ready” until all four are answered. ", 14.5, WHITE, True),
    ("This sequence is the binary definition of done.", 14.5, AMBER, True)]], line_spacing=1.1)
footer(s, "Case study · the gate")

# ---- CS4: the three core docs ---------------------------------------------
s = slide()
kicker(s, "Case study · what got built"); title(s, "Three docs do the load-bearing work")
docs3=[
    ("contracts.md","WHAT MUST NOT BREAK","Machine-checkable contract for all 34 /api routes: method, request/response shape, status codes, error envelope, auth class. Built FROM CODE, stamped with the commit SHA so staleness is detectable.", TEAL),
    ("invariants.md","THE “CHANGE-X-THEN-Y” MAP","The #1 cause of silent breakage: changing one side of a coupling and forgetting the other — especially across repos. Imperative checklists, some BLOCK-class: auth traps, FE↔automation links, data contracts, secrets.", AMBER),
    ("verification.md","HOW TO PROVE IT'S SAFE","Per change-type, the exact check that proves safety — and honest about what can't be proven locally: green=static check, yellow=CI, red=needs a deployed lane, human=judgment. “An app build passing is NOT enough.”", GREEN),
]
y=2.05; rh=1.45
for k,(name,tag,d,col) in enumerate(docs3):
    cy=y+k*(rh+0.12)
    box(s, 0.9, cy, 11.55, rh, fill=PANEL)
    box(s, 0.9, cy, 0.12, rh, fill=col, shape=MSO_SHAPE.RECTANGLE)
    text(s, 1.3, cy+0.14, 3.1, rh-0.25, [[(name, 16.5, WHITE, True)],[(tag, 10.5, col, True)]], space_after=4, line_spacing=1.0)
    text(s, 4.5, cy+0.14, 7.7, rh-0.25, [[(d, 12.5, MUTE)]], anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.08)
footer(s, "Case study · the docs")

# ---- CS5: supporting docs --------------------------------------------------
s = slide()
kicker(s, "Case study · the supporting cast"); title(s, "And the docs that keep it honest")
supp=[
    ("features-map.md","Per-feature blast radius — every feature's pages, routes, permissions, upstreams, cache tags, plus its QA page-objects & tests, in one place."),
    ("contract-drift-guard","Authority order (code > types > doc) + SHA staleness check, so contracts.md never silently rots."),
    ("blocked-items.md","Facts only humans can supply (compliance regime, on-call, owners) — explicit BLOCKERS, never invented."),
    ("runbooks/ (×6)","On-call playbooks: upstream down, Databricks pool, cache, auth/Okta, secrets, rollback."),
    ("ADR-0004","Records WHY this approach was chosen — the question a future engineer will ask."),
    ("Reviewed","9 specialist agents + 7 strategic lenses signed off before any of it landed."),
]
for k,(h,d) in enumerate(supp):
    cx=0.9+(k%2)*5.9; cy=2.1+(k//3)*0 + (k//2)*1.45
    box(s, cx, cy, 5.65, 1.3, fill=PANEL)
    box(s, cx, cy, 0.1, 1.3, fill=TEAL if k%2==0 else AMBER, shape=MSO_SHAPE.RECTANGLE)
    text(s, cx+0.3, cy+0.16, 5.2, 1.05, [[(h, 14.5, WHITE, True)],[(d, 11.5, MUTE)]], space_after=4, line_spacing=1.05)
footer(s, "Case study · supporting docs")

# ---- CS6: guardrails -------------------------------------------------------
s = slide()
kicker(s, "Case study · where autonomy stops"); title(s, "The docs fence off what agents can't self-approve")
box(s, 0.9, 2.2, 5.7, 2.7, fill=PANEL); box(s, 0.9, 2.2, 5.7, 0.12, fill=RED, shape=MSO_SHAPE.RECTANGLE)
text(s, 1.2, 2.45, 5.1, 2.3, [[("PII-bearing routes", 16, WHITE, True)],
    [("(person lookup, etc.)", 12, MUTE)],[("",5,WHITE)],
    [("→ blocked until the compliance / privacy regime is named by an owner.", 13, AMBER)]], space_after=6, line_spacing=1.12)
box(s, 6.75, 2.2, 5.7, 2.7, fill=PANEL); box(s, 6.75, 2.2, 5.7, 0.12, fill=RED, shape=MSO_SHAPE.RECTANGLE)
text(s, 7.05, 2.45, 5.1, 2.3, [[("Known fail-open auth routes", 16, WHITE, True)],
    [("the ones the review found", 12, MUTE)],[("",5,WHITE)],
    [("→ blocked until the security fix (tracked as T-004) lands.", 13, AMBER)]], space_after=6, line_spacing=1.12)
text(s, 0.9, 5.3, 11.5, 0.9, [[("“We don't know” is written down as a BLOCKER with an owner — ", 15, WHITE, True),
    ("not guessed, not silently assumed safe.", 15, AMBER, True, True)]], line_spacing=1.15)
footer(s, "Case study · guardrails")

# ---- CS7: PROOF — found real bugs (the money slide) -----------------------
s = slide()
kicker(s, "Case study · proof it works"); title(s, "Writing the contracts found real production holes")
text(s, 0.9, 2.0, 11.5, 0.7, [[("Building contracts.md ", 14.5, WHITE),("from code", 14.5, AMBER, True),
    (" mechanically surfaced two authorization holes the prose docs had missed:", 14.5, WHITE)]], line_spacing=1.1)
bugs=[
    ("/api/contact-validation/:id","A MUTATING endpoint (flips PRP feature flags) reachable by ANY authenticated user, regardless of role."),
    ("/api/databricks/person-lookup","A PII query surface — fail-open."),
]
for k,(r,d) in enumerate(bugs):
    cy=2.85+k*1.25
    box(s, 0.9, cy, 11.55, 1.1, fill=PANEL)
    box(s, 0.9, cy, 0.12, 1.1, fill=RED, shape=MSO_SHAPE.RECTANGLE)
    text(s, 1.3, cy+0.16, 11.0, 0.85, [[(r, 15.5, RED, True, False, "Consolas")],[(d, 13, WHITE)]], space_after=4, line_spacing=1.08)
box(s, 0.9, 5.55, 11.55, 1.0, fill=TEALDK)
text(s, 1.2, 5.68, 11.0, 0.8, [[("The docs doing their job: writing the contract surfaced real holes. ", 14.5, WHITE, True),
    ("Both are now documented + tracked for fix — not silently shipped.", 14.5, WHITE)]], line_spacing=1.12)
footer(s, "Case study · proof", accent=AMBER)

# ---- CS8: before vs after table -------------------------------------------
s = slide()
kicker(s, "Case study · before vs. after"); title(s, "What changed")
rows=[
    ("API contracts","prose only","34 routes, code-sourced"),
    ("Coupling / blast radius","none","invariants + features-map"),
    ("Proof-of-safety","implicit","per-change playbook"),
    ("Cross-repo awareness","none","FE ↔ QA links mapped"),
    ("Staleness detection","none","SHA stamp + drift guard"),
    ("Unknown facts","scattered / implied","explicit owned blockers"),
]
# header
hy=2.05
text(s, 3.0, hy, 3.0, 0.4, [[("", 12, MUTE)]])
text(s, 6.15, hy, 3.0, 0.4, [[("BEFORE (T-001)", 12.5, RED, True)]])
text(s, 9.35, hy, 3.0, 0.4, [[("AFTER (T-002)", 12.5, GREEN, True)]])
for k,(label,b,a) in enumerate(rows):
    cy=2.5+k*0.62
    if k%2==0: box(s, 0.9, cy, 11.55, 0.58, fill=PANEL)
    text(s, 1.1, cy, 4.9, 0.58, [[(label, 13, WHITE, True)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 6.15, cy, 3.0, 0.58, [[(b, 12.5, MUTE)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, 9.35, cy, 3.1, 0.58, [[(a, 12.5, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
box(s, 0.9, 6.25, 11.55, 0.62, fill=TEALDK)
text(s, 1.1, 6.25, 11.2, 0.62, [[("Net effect:   ", 13.5, WHITE, True),("“describes it”   →   ", 13.5, MUTE),
    ("“safe to act on it”  (under human gate)", 13.5, AMBER, True)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s, "Case study · before/after")

# ---- CS9: bottom line ------------------------------------------------------
s = slide()
section(s, "Case study · bottom line",
        "“If I change this, what breaks — and how do I prove I didn't?”",
        "The docs now answer that for any change. That is the difference between docs that "
        "DESCRIBE a system and docs that are SAFE TO ACT ON. The codebase was never touched — "
        "the living memory simply became trustworthy enough to supervise real work against.",
        accent=AMBER)

# ===========================================================================
# 15 — RECAP
# ===========================================================================
s = slide()
kicker(s, "Recap"); title(s, "What you get")
recap = [
    "Every agent starts grounded in the same living memory — no cold starts, no invented facts.",
    "contracts.md + the regression-safety gate mean changes are made with eyes open to blast radius.",
    "9 specialists + 3 rounds + 7 lenses catch what a single pass misses.",
    "Hooks + CI + branch protection make the conventions mechanical, not aspirational.",
    "One framework, single- or multi-repo, solo founder to enterprise team.",
]
for k,t in enumerate(recap):
    cy=2.05+k*0.92
    box(s, 0.9, cy, 11.55, 0.78, fill=PANEL)
    box(s, 1.15, cy+0.24, 0.3, 0.3, fill=TEAL, shape=MSO_SHAPE.OVAL)
    text(s, 1.7, cy, 10.5, 0.78, [[(t, 14.5, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s, "Recap")

# ===========================================================================
# 16 — GET STARTED / CLOSE
# ===========================================================================
s = slide()
box(s, 0, 0, 13.333, 7.5, fill=INK, shape=MSO_SHAPE.RECTANGLE)
box(s, -1.3, 4.6, 5, 5, fill=PANEL, shape=MSO_SHAPE.OVAL)
box(s, 10.6, -1.4, 4.2, 4.2, fill=TEALDK, shape=MSO_SHAPE.OVAL)
accent_bar(s, 0.95, 2.2, w=2.0)
text(s, 0.9, 2.4, 11.5, 2.0, [
    [("Start in five minutes", 44, WHITE, True, False, FONTH)],
    [("Fork it. Replace the placeholders. Install the hooks. Let the agent build /docs.", 18, TEAL)],
], space_after=12, line_spacing=1.1)
box(s, 0.9, 4.5, 11.55, 1.4, fill=PANEL)
text(s, 1.25, 4.66, 11.0, 1.1, [
    [("Read first:  ", 14, MUTE),("ADOPTING.md", 14, AMBER, True),("  (fork → enforced)        ", 14, MUTE),
     ("AGENTS.md", 14, AMBER, True),("  (the canonical conventions)", 14, MUTE)],
    [("Then just start your agent — it reads the rules and bootstraps the rest.", 13.5, WHITE)],
], space_after=8, line_spacing=1.2)
text(s, 0.9, 6.5, 11.5, 0.5, [[("Agentic Starter  —  conventions + memory + enforcement, so AI you can trust.", 13, MUTE, False, True)]])

# ---- save ------------------------------------------------------------------
assert SLIDE_NO == TOTAL, f"TOTAL={TOTAL} but rendered {SLIDE_NO} slides — fix TOTAL"
out = "Agentic-Starter-Overview.pptx"
prs.save(out)
print("wrote", out, "slides:", len(prs.slides._sldIdLst))
